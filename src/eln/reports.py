"""Office-document report generators: Word (.docx), Excel (.xlsx), charts (.png).

Ported from Eyal's lab_team scientist tools, adapted to this package. These are
pure, deterministic, fully local generators — they never call an LLM and nothing
leaves the workstation, so they sit outside the Sanitizer/IP boundary.

matplotlib must select a non-interactive backend BEFORE pyplot is imported
because Streamlit and the CLI both run headless.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
REPORTS_DIR = ROOT / "artifacts" / "reports"


def _out(name: str, suffix: str) -> Path:
    """Resolve a SHORT filename (no extension, no directory) to a path under
    REPORTS_DIR, stripping anything that is not alphanumeric/-/_ to defend
    against path traversal and full-path arguments from an LLM tool call."""
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in name) or "output"
    return REPORTS_DIR / f"{safe}{suffix}"


# ── Charts ────────────────────────────────────────────────────

def plot_chart(
    chart_type: str,
    labels: list[str],
    values: list[float],
    title: str,
    output_name: str,
) -> Path:
    """Render a 'bar' or 'line' chart from labels + numeric values to a PNG.
    Returns the absolute path to the saved image."""
    if len(labels) != len(values):
        raise ValueError(f"labels ({len(labels)}) and values ({len(values)}) must match")
    vals = [float(v) for v in values]
    fig, ax = plt.subplots(figsize=(8, 4.5))
    if chart_type == "line":
        ax.plot(labels, vals, marker="o")
    else:
        ax.bar(labels, vals)
    ax.set_title(title)
    fig.autofmt_xdate(rotation=30)
    fig.tight_layout()
    path = _out(output_name, ".png")
    fig.savefig(path, dpi=120)
    plt.close(fig)
    return path


# ── Excel ─────────────────────────────────────────────────────

def to_excel(
    sheet_name: str,
    headers: list[str],
    rows: list[list[Any]],
    output_name: str,
) -> Path:
    """Write a table (headers + rows) to an .xlsx workbook. Returns the path."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Sheet1")[:31]
    if headers:
        ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    path = _out(output_name, ".xlsx")
    wb.save(path)
    return path


# ── Word ──────────────────────────────────────────────────────

def to_word(title: str, sections: list[dict], output_name: str) -> Path:
    """Write a Word (.docx) report. `sections` is a list of dicts, each with a
    'heading' and optional 'text', 'table' (list of rows, first row = header),
    and 'image' (absolute path to a PNG, e.g. from plot_chart). Returns the path."""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(title, level=0)
    for sec in sections or []:
        if sec.get("heading"):
            doc.add_heading(str(sec["heading"]), level=1)
        if sec.get("text"):
            doc.add_paragraph(str(sec["text"]))
        table = sec.get("table")
        if table:
            rows = list(table)
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=cols)
            t.style = "Light Grid Accent 1"
            for r in rows:
                cells = t.add_row().cells
                for i, val in enumerate(r):
                    cells[i].text = "" if val is None else str(val)
        image = sec.get("image")
        if image and Path(image).exists():
            doc.add_picture(str(image), width=Inches(6))
    path = _out(output_name, ".docx")
    doc.save(path)
    return path


# ── High-level assembler for the Flow ─────────────────────────

def experiment_report(
    record: dict,
    prediction: dict,
    report_text: str,
    safety: dict | None = None,
    output_name: str = "experiment-report",
) -> Path:
    """Assemble a complete Word report from one Flow run: the structured record,
    the model prediction (with a yield chart), the sanitised report text, and any
    safety alert. Returns the path to the .docx.

    Inputs come straight from ELNState, so this stays inside the IP boundary the
    Flow already enforces — report_text is the sanitised text from write_report.
    """
    sections: list[dict] = []

    safety = safety or {}
    if safety.get("severity") == "critical":
        sections.append({
            "heading": "⛔ Safety alert",
            "text": "Critical reagent incompatibility detected — review the Safety "
                    "check before any bench work.",
        })

    if record:
        sections.append({
            "heading": "Experiment record",
            "table": [["Parameter", "Value"]] + [[k, v] for k, v in record.items()],
        })

    chart_path: Path | None = None
    if prediction:
        # bool is a subclass of int — exclude it so a True/False flag (e.g.
        # predicted_qc_pass) is not charted as a meaningless 1.0 bar.
        numeric = {
            k: v for k, v in prediction.items()
            if isinstance(v, (int, float)) and not isinstance(v, bool)
        }
        if numeric:
            chart_path = plot_chart(
                "bar", list(numeric.keys()), list(numeric.values()),
                "Predicted outcome", f"{output_name}-prediction",
            )
        sections.append({
            "heading": "Prediction (scikit-learn)",
            "table": [["Metric", "Value"]] + [[k, v] for k, v in prediction.items()],
            **({"image": str(chart_path)} if chart_path else {}),
        })

    if report_text:
        sections.append({"heading": "Report", "text": report_text})

    return to_word("ELN Experiment Report", sections, output_name)
