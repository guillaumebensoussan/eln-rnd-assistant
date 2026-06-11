"""Scientist agent tools — DB reads, RAG literature, and Word/Excel/chart output.

matplotlib must use a non-interactive backend (set BEFORE importing pyplot)
because this runs headless in a CLI.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

from crewai.tools import tool  # noqa: E402

from lab_team import paths  # noqa: F401,E402  (extends sys.path) — must be first
import eln_db  # noqa: E402

from lab_team.config import RESULTS_DIR, rag_config  # noqa: E402
from lab_team.tools._shared import as_json, tool_guard  # noqa: E402


def _out(name: str, suffix: str) -> Path:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in name) or "output"
    return RESULTS_DIR / f"{safe}{suffix}"


# ── Data sources ──────────────────────────────────────────────

@tool("read_project_data")
def read_project_data(project_code: str) -> str:
    """Read everything about a project for reporting: the project record, each
    experiment (with yields and status), and the materials used in each
    experiment. Returns structured JSON."""

    def _read(code: str) -> dict:
        summary = eln_db.get_project_summary(code)
        for exp in summary["experiments"]:
            exp["materials"] = eln_db.list_experiment_materials(exp["code"])
        return summary

    return tool_guard(_read)(project_code)


@tool("get_synthesis_facts")
def get_synthesis_facts(compound: str, pdf_path: str) -> str:
    """Extract structured synthesis facts (starting materials, conditions,
    workup, yield, characterization) for a compound from a chemistry PDF, using
    the RAG pipeline. `pdf_path` is a local PDF path or URL. Runs in-process;
    first run downloads the embedding model."""

    def _facts(compound: str, pdf_path: str) -> dict:
        from pipeline.orchestrator import run_pipeline

        result = run_pipeline(
            source=pdf_path,
            query_compound=compound,
            output_dir=str(RESULTS_DIR),
            config=rag_config(),
        )
        return result.get("synthesis", result)

    return tool_guard(_facts)(compound, pdf_path)


# ── Output generators ─────────────────────────────────────────

@tool("plot_chart")
def plot_chart(
    chart_type: str,
    labels: list[str],
    values: list[float],
    title: str,
    output_name: str,
) -> str:
    """Render a 'bar' or 'line' chart from labels + numeric values to a PNG.
    Returns the absolute path to the saved image (embed it in a Word report).
    `output_name` must be a SHORT filename WITHOUT extension and WITHOUT any
    directory path, e.g. 'yield-comparison' (not a full path)."""

    def _plot(chart_type: str, labels: list, values: list, title: str, output_name: str) -> str:
        if len(labels) != len(values):
            raise ValueError(f"labels ({len(labels)}) and values ({len(values)}) must match")
        vals = [float(v) for v in values]
        fig, ax = plt.subplots(figsize=(8, 4.5))
        if chart_type == "line":
            ax.plot(labels, vals, marker="o")
        else:
            ax.bar(labels, vals)
        ax.set_title(title)
        fig.autofmt_xdate(rotation=30)
        fig.tight_layout()
        path = _out(output_name, ".png")
        fig.savefig(path, dpi=120)
        plt.close(fig)
        return f"Chart saved: {path}"

    return tool_guard(_plot)(chart_type, labels, values, title, output_name)


@tool("generate_excel_report")
def generate_excel_report(
    sheet_name: str,
    headers: list[str],
    rows: list[list[Any]],
    output_name: str,
) -> str:
    """Write a table (headers + rows) to an .xlsx workbook. Returns the
    absolute path to the saved file."""

    def _excel(sheet_name: str, headers: list, rows: list, output_name: str) -> str:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = (sheet_name or "Sheet1")[:31]
        if headers:
            ws.append(list(headers))
        for row in rows:
            ws.append(list(row))
        path = _out(output_name, ".xlsx")
        wb.save(path)
        return f"Excel report saved: {path}"

    return tool_guard(_excel)(sheet_name, headers, rows, output_name)


@tool("generate_word_report")
def generate_word_report(title: str, sections: list[dict], output_name: str) -> str:
    """Write a Word (.docx) report. `sections` is a list of dicts, each with a
    'heading' and optional 'text', 'table' (list of rows, first row = header),
    and 'image' (absolute path to a PNG, e.g. from plot_chart). Returns the
    absolute path to the saved document.
    `output_name` must be a SHORT filename WITHOUT extension and WITHOUT any
    directory path, e.g. 'Suzuki-Optimization-Report' (not a full path)."""

    def _word(title: str, sections: list, output_name: str) -> str:
        from docx import Document
        from docx.shared import Inches

        doc = Document()
        doc.add_heading(title, level=0)
        for sec in sections or []:
            if sec.get("heading"):
                doc.add_heading(str(sec["heading"]), level=1)
            if sec.get("text"):
                doc.add_paragraph(str(sec["text"]))
            table = sec.get("table")
            if table:
                rows = list(table)
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=cols)
                t.style = "Light Grid Accent 1"
                for r in rows:
                    cells = t.add_row().cells
                    for i, val in enumerate(r):
                        cells[i].text = "" if val is None else str(val)
            image = sec.get("image")
            if image and Path(image).exists():
                doc.add_picture(str(image), width=Inches(6))
        path = _out(output_name, ".docx")
        doc.save(path)
        return f"Word report saved: {path}"

    return tool_guard(_word)(title, sections, output_name)


SCIENTIST_TOOLS = [
    read_project_data,
    get_synthesis_facts,
    plot_chart,
    generate_excel_report,
    generate_word_report,
]
